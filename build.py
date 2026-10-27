#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build.py — 產生「Moscito 叫序資料庫」靜態網站（HCP 版）

架構照搬 bridge-relay-bidding（北歐接力精確體系網站）：
  原始資料/*.xlsx  ->  網站/叫序資料/*.html，外加 網站/index.html

重新執行本腳本即可依來源檔案的最新內容重新產生整個 網站/ 目錄的 .html。
網站/assets/style.css、site.js 為手動維護，不會被覆寫。
"""
import os
import re
import html
import shutil

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "原始資料")
OUT = os.path.join(ROOT, "網站")

SUIT_RED = set("♥♦")
SUIT_BLACK = set("♠♣")


# --------------------------------------------------------------------------
# 共用版型
# --------------------------------------------------------------------------

SITE_NAME = "Moscito 叫序資料庫"
SOURCE_BOOK = "《Terrorist's Moscito》（Bo-Yin Yang, 2000）"

# (xlsx 檔名, 頁面標題, 首頁卡片角標, 首頁／頁首簡介)
PAGES = [
    ("一梅花開叫.xlsx", "一梅花開叫", "Clubs",
     "1♣（約 16+ 點強梅花）的應叫、1♣–1♦ 後續、第二接力、積極應叫、特殊應叫與 1♣ 被干擾。"),
    ("一方塊開叫.xlsx", "一方塊開叫", "Diamonds",
     "1♦（約 10-14 點，否認四張高花）的應叫、1♦–1NT 接力、1♦–1♥ 等待接力與被干擾。"),
    ("高花開叫.xlsx", "一高花開叫", "Majors",
     "1♥／1♠（約 9-14 點，canapé）的應叫、加叫總表、Stenberg 2NT、接力與被干擾。"),
    ("一無王開叫.xlsx", "一無王開叫", "No Trump",
     "1NT（約 9-14 點，兩高花各 4+）、2♣ 接力、第三四家 1♦，以及自然無王體系。"),
    ("阻擊開叫.xlsx", "阻擊開叫", "Preempts",
     "二階弱開叫的接力、弱二方案、三階以上與強阻擊。"),
    ("接力方案.xlsx", "接力方案", "Relays",
     "第五章：單套、兩三套、均型方案，三套速查、強度範圍、螺旋掃描與斷鏈。"),
    ("競叫.xlsx", "競叫", "Competition",
     "第三、四章：防守叫牌、Asptro、trap 防守、Rubensohl、Transfer lebensohl 與強制性 pass。"),
]

NAV_ITEMS = [(t, "叫序資料/%s.html" % t) for _, t, _, _ in PAGES]


def suit_colorize(text):
    """把文字中的 ♥♦ 轉紅、♠♣ 轉黑（沿用速查頁的 .r/.b 慣例），並保留換行。"""
    out = []
    for ch in text:
        if ch in SUIT_RED:
            out.append('<span class="s r">%s</span>' % ch)
        elif ch in SUIT_BLACK:
            out.append('<span class="s b">%s</span>' % ch)
        elif ch == "\n":
            out.append("<br>")
        else:
            out.append(html.escape(ch))
    return "".join(out)


def render_nav(root_prefix, current_href=None):
    parts = []
    parts.append('<nav class="site-nav">')
    parts.append('<div class="nav-inner">')
    parts.append('<a class="nav-brand" href="%sindex.html">%s</a>' % (root_prefix, SITE_NAME))
    parts.append('<button class="nav-toggle" aria-label="開啟選單" type="button">'
                  '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">'
                  '<path d="M3 6h18M3 12h18M3 18h18"/></svg></button>')
    parts.append('<ul class="nav-links" role="menubar">')
    for label, href in NAV_ITEMS:
        active = " is-active" if href == current_href else ""
        parts.append('<li class="nav-item" role="none"><a class="nav-link%s" href="%s%s">%s</a></li>' %
                      (active, root_prefix, href, label))
    parts.append('</ul>')
    parts.append('</div></nav>')
    return "".join(parts)


PAGE_TEMPLATE = """<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{description}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@500;700&family=Noto+Sans+TC:wght@400;500;700&family=Noto+Serif+TC:wght@600;700&display=swap">
<link rel="stylesheet" href="{root}assets/style.css">
</head>
<body>
{nav}
<div class="wrap">
{body}
</div>
<script src="{root}assets/site.js"></script>
</body>
</html>
"""


def render_page(title, description, root_prefix, current_href, body_html):
    return PAGE_TEMPLATE.format(
        title=html.escape(title),
        description=html.escape(description),
        root=root_prefix,
        nav=render_nav(root_prefix, current_href),
        body=body_html,
    )


def page_head(eyebrow, h1, meta_lines, lede_html, crumb_html=""):
    meta = "<br>".join(meta_lines)
    return """
{crumb}
<header class="page-head">
  <div class="mast">
    <div>
      <div class="eyebrow">{eyebrow}</div>
      <h1>{h1}</h1>
    </div>
    <div class="mast-meta">{meta}</div>
  </div>
  <p class="lede">{lede}</p>
</header>
""".format(crumb=crumb_html, eyebrow=eyebrow, h1=h1, meta=meta, lede=lede_html)


def write_file(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    print("wrote", os.path.relpath(path, ROOT))


# --------------------------------------------------------------------------
# xlsx -> HTML（叫序資料）
#
# 每個工作表先被 analyze_sheet() 拆成「前導說明列／表頭／資料列／註腳」，
# 再依表頭形狀分派給對應的語意化 renderer（沿用網站既有的排版元件，
# 不再輸出 Excel 原始底色）。資料列中形如「→ 表4「A分支」」「見表 9」的
# 文字，會被 linkify_refs() 轉成可點擊的分頁跳轉連結。
# --------------------------------------------------------------------------

HEADER_NAVY = "1F3864"


def is_header_style(cell):
    """判斷這個儲存格是否具備原始表格「表頭列」的樣式（粗體＋深藍底），
    用來定位表頭列，以及「說明」表中重複出現的小節分隔列。"""
    if not (cell.font and cell.font.bold):
        return False
    fill = cell.fill
    if not (fill and fill.fill_type == "solid" and fill.fgColor):
        return False
    fg = fill.fgColor
    if fg.type != "rgb" or not fg.rgb:
        return False
    return fg.rgb[-6:].upper() == HEADER_NAVY


def cell_text(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None:
        return ""
    v = str(v)
    if v.startswith("'"):
        # Excel「強制文字」前綴的殘留單引號，不是內容的一部分
        v = v[1:]
    return v


def analyze_sheet(ws):
    max_col = ws.max_column
    max_row = ws.max_row

    header_row = None
    for r in range(1, max_row + 1):
        if is_header_style(ws.cell(row=r, column=1)):
            header_row = r
            break
    if header_row is None:
        header_row = 1

    leading = [cell_text(ws, r, 1).strip() for r in range(1, header_row)]
    leading = [t for t in leading if t]

    header = [cell_text(ws, header_row, c).strip() for c in range(1, max_col + 1)]
    while header and header[-1] == "":
        header.pop()

    rows = []
    for r in range(header_row + 1, max_row + 1):
        vals = [cell_text(ws, r, c) for c in range(1, max_col + 1)]
        if not any(v.strip() for v in vals):
            continue
        if is_header_style(ws.cell(row=r, column=1)):
            rows.append(("SUB", vals[0].strip()))
            continue
        rows.append(("DATA", vals))

    footnotes = []
    while rows and rows[-1][0] == "DATA":
        vals = rows[-1][1]
        nonempty = [i for i, v in enumerate(vals) if v.strip()]
        if nonempty == [0]:
            footnotes.insert(0, vals[0].strip())
            rows.pop()
        else:
            break

    return dict(leading=leading, header=header, rows=rows, footnotes=footnotes, max_col=max_col)


BID_TOKEN_RE = re.compile(r'(\d)(<span class="s [rb]">[♠♥♦♣]</span>|NT)')


def bidify(html_text):
    return BID_TOKEN_RE.sub(lambda m: '<span class="bid">%s</span>' % m.group(0), html_text)


BRANCH_CODE_RE = re.compile(r"^([A-Za-z])(?:[\s　]|$)")


def extract_branch_code(raw):
    if not raw:
        return None
    m = BRANCH_CODE_RE.match(raw.strip())
    return m.group(1) if m else None


REF_RE = re.compile(r"表\s*(\d+(?:[、,]\s*\d+)*)(?![\dA-Za-z])(?:「([^」]*)」)?")


def linkify_refs(raw, num_to_index, merged_indices, current_code):
    """把「→ 表4「A分支」」「見表 9」「見表 3、4」這類文字轉成可點擊的分頁跳轉連結。"""

    def repl(m):
        nums_str, label = m.group(1), m.group(2)
        nums = re.split(r"[、,]\s*", nums_str)
        links = []
        for ns in nums:
            idx = num_to_index.get(int(ns))
            if idx is None:
                links.append(ns)
                continue
            dom = "sheet-%d" % (idx + 1)
            attrs = 'class="jump-link" data-sheet="%s"' % dom
            if idx in merged_indices and current_code:
                attrs += ' data-anchor="branch-%d-%s"' % (idx + 1, current_code)
            links.append("<a %s>%s</a>" % (attrs, ns))
        suffix = ("「%s」" % label) if label else ""
        return "表" + "、".join(links) + suffix

    return REF_RE.sub(repl, raw)


# 跨頁連結：資料裡寫「見接力方案檔」「見競叫檔 Rubensohl」這類文字時，連到對應頁面
PAGE_REF_ALIASES = {
    "1♣ 開叫檔": "一梅花開叫", "1♦ 開叫檔": "一方塊開叫", "高花開叫檔": "一高花開叫",
    "1NT 開叫檔": "一無王開叫", "阻擊開叫檔": "阻擊開叫", "接力方案檔": "接力方案",
    "《接力方案》檔": "接力方案", "《競叫》檔": "競叫", "競叫檔": "競叫",
}
PAGE_REF_RE = re.compile("|".join(re.escape(k) for k in sorted(PAGE_REF_ALIASES, key=len, reverse=True)))


def linkify_pages(raw):
    def repl(m):
        return '<a class="tool-link" href="%s.html">%s</a>' % (PAGE_REF_ALIASES[m.group(0)], m.group(0))
    return PAGE_REF_RE.sub(repl, raw)


def render_cell(raw, num_to_index, merged_indices, current_code=None):
    """統一的儲存格內容渲染：跳轉連結 → 花色上色／逃逸 → 叫品 mono 樣式。"""
    if raw is None:
        raw = ""
    linked = linkify_refs(raw, num_to_index, merged_indices, current_code)
    linked = linkify_pages(linked)
    parts = re.split(r"(<[^>]+>)", linked)
    out = []
    for part in parts:
        if part.startswith("<"):
            out.append(part)
            continue
        seg = suit_colorize(part)
        seg = bidify(seg)
        out.append(seg)
    return "".join(out)


def render_lead_block(leading, num_to_index, merged_indices):
    if not leading:
        return ""
    parts = ['<div class="lead-block">']
    parts.append('<div class="lead-title">%s</div>' %
                  render_cell(leading[0], num_to_index, merged_indices))
    if len(leading) > 1:
        parts.append('<div class="lead-context">%s</div>' %
                      render_cell(leading[1], num_to_index, merged_indices))
    for extra in leading[2:]:
        parts.append('<div class="lead-def">%s</div>' %
                      render_cell(extra, num_to_index, merged_indices))
    parts.append("</div>")
    return "".join(parts)


def render_footnotes(footnotes, num_to_index, merged_indices):
    if not footnotes:
        return ""
    lines = [render_cell(f, num_to_index, merged_indices) for f in footnotes]
    return '<div class="trap">%s</div>' % "<br>".join(lines)


PARENT_ROW_RE = re.compile(r"^▸\s*")
CHILD_ROW_RE = re.compile(r"^　+")
CHILD_NOTE_RE = re.compile(r"^(以下\s*\d+\s*列為子接力|細分見下\s*[一二兩三四五六七八九十\d]*\s*列)$")


def render_ladder_row(vals, n, code_col, num_to_index, merged_indices, toggle_group=None, child_of=None):
    """渲染單一列 <tr>。toggle_group：這列是子接力的父列，附上展開／收合按鈕。
    child_of：這列是子接力，預設收合、與父列共用同一個 group 供 JS 對應。"""
    current_code = extract_branch_code(vals[0])
    tds = []
    for ci, raw in enumerate(vals):
        if ci == 0 and code_col:
            content = ('<span class="code-badge">%s</span>' % html.escape(raw.strip())
                        if raw.strip() else "&nbsp;")
        elif ci == 0 and current_code and len(raw.strip()) > 1:
            rest = raw.strip()[1:].strip()
            content = '<span class="code-badge">%s</span> %s' % (
                html.escape(current_code), render_cell(rest, num_to_index, merged_indices, current_code))
        else:
            content = render_cell(raw, num_to_index, merged_indices, current_code)
        if ci == 0 and toggle_group:
            content = ('<button class="ladder-toggle" type="button" aria-expanded="false">'
                       '<span class="tri">▸</span></button>%s' % content)
        if not content:
            content = "&nbsp;"
        cls = ' class="cell-muted"' if ci == n - 1 else ""
        tds.append("<td%s>%s</td>" % (cls, content))
    if toggle_group:
        return '<tr class="ladder-parent" data-group="%s">%s</tr>' % (toggle_group, "".join(tds))
    if child_of:
        return '<tr class="ladder-child" data-group="%s" hidden>%s</tr>' % (child_of, "".join(tds))
    return "<tr>%s</tr>" % "".join(tds)


def render_ladder(analysis, sheet_index, num_to_index, merged_indices):
    header = analysis["header"]
    n = len(header)
    code_col = bool(header) and header[0] == "代號"

    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    ths = "".join("<th>%s</th>" % html.escape(h) for h in header)
    trs = []
    rows = analysis["rows"]
    # 巢狀標記（▸ 上層／全形空白子接力）通常在第一欄；
    # 第一欄是「序列」的表，標記會放在第二欄（叫品欄）
    mc = 0
    data_rows = [p for k, p in rows if k == "DATA"]
    if not any(PARENT_ROW_RE.match(p[0] or "") for p in data_rows) and \
            any(len(p) > 1 and PARENT_ROW_RE.match(p[1] or "") for p in data_rows):
        mc = 1
    group_seq = 0
    i = 0
    while i < len(rows):
        kind, payload = rows[i]
        if kind == "SUB":
            trs.append('<tr class="subhead"><td colspan="%d">%s</td></tr>' %
                        (n, render_cell(payload, num_to_index, merged_indices)))
            i += 1
            continue

        vals = (payload + [""] * n)[:n]
        if vals[0].lstrip().startswith("＊") and not any(v.strip() for v in vals[1:]):
            trs.append('<tr class="note-row"><td colspan="%d">%s</td></tr>' %
                        (n, render_cell(vals[0], num_to_index, merged_indices)))
            i += 1
            continue
        children = []
        j = i + 1
        while j < len(rows) and rows[j][0] == "DATA":
            cvals = (rows[j][1] + [""] * n)[:n]
            if not CHILD_ROW_RE.match(cvals[mc]):
                break
            children.append(cvals)
            j += 1

        if children:
            group_seq += 1
            group = "ladder-group-%d-%d" % (sheet_index + 1, group_seq)
            vals[mc] = PARENT_ROW_RE.sub("", vals[mc])
            if CHILD_NOTE_RE.match(vals[-1].strip()):
                vals[-1] = ""
            trs.append(render_ladder_row(vals, n, code_col, num_to_index, merged_indices, toggle_group=group))
            for cvals in children:
                cvals[mc] = CHILD_ROW_RE.sub("", cvals[mc])
                trs.append(render_ladder_row(cvals, n, code_col, num_to_index, merged_indices, child_of=group))
            i = j
        else:
            vals[mc] = PARENT_ROW_RE.sub("", vals[mc])
            trs.append(render_ladder_row(vals, n, code_col, num_to_index, merged_indices))
            i += 1

    parts.append('<div class="table-scroll"><table class="bidtbl"><thead><tr>%s</tr></thead>'
                  '<tbody>%s</tbody></table></div>' % (ths, "".join(trs)))
    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_merged_branch(analysis, sheet_index, num_to_index, merged_indices):
    header = analysis["header"]
    inner_headers = header[2:]

    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * len(header))[:len(header)]
        code = vals[0].strip()
        context = vals[1]
        rest = vals[2:]
        if not groups or groups[-1][0] != code:
            groups.append([code, context, [rest]])
        else:
            groups[-1][2].append(rest)

    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    for code, context, rows in groups:
        anchor_id = "branch-%d-%s" % (sheet_index + 1, code)
        parts.append('<div class="branch-group" id="%s">' % anchor_id)
        parts.append('<div class="branch-head"><span class="code-badge">%s</span><span class="ctx">%s</span></div>' %
                      (html.escape(code), render_cell(context, num_to_index, merged_indices)))
        ths = "".join("<th>%s</th>" % html.escape(h) for h in inner_headers)
        trs = []
        for rvals in rows:
            tds = []
            for ci, raw in enumerate(rvals):
                content = render_cell(raw, num_to_index, merged_indices, code) or "&nbsp;"
                cls = ' class="cell-muted"' if ci == len(rvals) - 1 else ""
                tds.append("<td%s>%s</td>" % (cls, content))
            trs.append("<tr>%s</tr>" % "".join(tds))
        parts.append('<div class="table-scroll"><table class="bidtbl"><thead><tr>%s</tr></thead>'
                      '<tbody>%s</tbody></table></div>' % (ths, "".join(trs)))
        parts.append("</div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_info(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    parts.append('<table class="deftable"><tbody>')
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            parts.append('</tbody></table><div class="deftable-sub">%s</div><table class="deftable"><tbody>' %
                          html.escape(payload))
            continue
        vals = (payload + ["", ""])[:2]
        label = render_cell(vals[0], num_to_index, merged_indices)
        desc = render_cell(vals[1], num_to_index, merged_indices)
        parts.append("<tr><td>%s</td><td>%s</td></tr>" % (label, desc))
    parts.append("</tbody></table>")
    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_hands(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * 10)[:10]
        exno = vals[0].strip()
        if not groups or groups[-1][0] != exno:
            groups.append([exno, []])
        groups[-1][1].append(vals)

    for exno, rows in groups:
        first = rows[0]
        source, situation = first[1], first[2]
        contract = next((r[8] for r in rows if r[8].strip()), "")
        comment = next((r[9] for r in rows if r[9].strip()), "")
        parts.append('<div class="example-card">')
        parts.append('<div class="example-head"><b>%s</b><span>%s</span><span>%s</span></div>' %
                      (html.escape(exno), render_cell(source, num_to_index, merged_indices),
                       render_cell(situation, num_to_index, merged_indices)))
        parts.append('<div class="example-body"><div class="table-scroll"><table class="handtbl"><tbody>')
        for r in rows:
            seat, sp, he, di, cl = r[3], r[4], r[5], r[6], r[7]
            parts.append(
                '<tr><td class="seat">%s</td>'
                '<td class="suit"><span class="s b">♠</span>%s</td>'
                '<td class="suit"><span class="s r">♥</span>%s</td>'
                '<td class="suit"><span class="s r">♦</span>%s</td>'
                '<td class="suit"><span class="s b">♣</span>%s</td></tr>' %
                (render_cell(seat, num_to_index, merged_indices),
                 render_cell(sp, num_to_index, merged_indices),
                 render_cell(he, num_to_index, merged_indices),
                 render_cell(di, num_to_index, merged_indices),
                 render_cell(cl, num_to_index, merged_indices)))
        parts.append("</tbody></table></div>")
        if contract.strip() or comment.strip():
            parts.append('<div class="example-result">')
            if contract.strip():
                parts.append('<span class="bid-lead">%s</span>' % render_cell(contract, num_to_index, merged_indices))
            if comment.strip():
                parts.append(render_cell(comment, num_to_index, merged_indices))
            parts.append("</div>")
        parts.append("</div></div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_auction(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * 6)[:6]
        exno = vals[0].strip()
        # 只有每個例子的第一列才有「例號」；後續列的例號欄位是空的，仍屬於同一組
        if exno and (not groups or groups[-1][0] != exno):
            groups.append([exno, vals[1], []])
        if not groups:
            groups.append(["", vals[1], []])
        groups[-1][2].append(vals[2:])

    for exno, situation, rows in groups:
        parts.append('<div class="example-card">')
        parts.append('<div class="example-head"><b>%s</b><span>%s</span></div>' %
                      (html.escape(exno), render_cell(situation, num_to_index, merged_indices)))
        parts.append('<div class="example-body"><div class="table-scroll"><table class="auctiontbl"><thead>'
                      '<tr><th>序</th><th>叫者</th><th>叫品</th><th>說明</th></tr></thead><tbody>')
        for seq, who, bid, note in rows:
            note_html = render_cell(note, num_to_index, merged_indices) if note.strip() else ""
            parts.append(
                '<tr><td class="seq">%s</td><td class="who">%s</td>'
                '<td>%s</td><td>%s</td></tr>' %
                (html.escape(seq.strip()), render_cell(who, num_to_index, merged_indices),
                 render_cell(bid, num_to_index, merged_indices), note_html))
        parts.append("</tbody></table></div></div></div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_sheet_panel(analysis, sheet_index, num_to_index, merged_indices):
    header = tuple(analysis["header"])
    if header[:2] == ("項目", "內容"):
        return render_info(analysis, num_to_index, merged_indices)
    if header[:2] == ("分支", "序列／前提"):
        return render_merged_branch(analysis, sheet_index, num_to_index, merged_indices)
    if header[:1] == ("例號",) and "方位" in header:
        return render_hands(analysis, num_to_index, merged_indices)
    if header[:1] == ("例號",) and "叫者" in header:
        return render_auction(analysis, num_to_index, merged_indices)
    return render_ladder(analysis, sheet_index, num_to_index, merged_indices)


def build_xlsx_page(xlsx_name, title, eyebrow, lede):
    xlsx_names = [xlsx_name] if isinstance(xlsx_name, str) else xlsx_name

    sheets = []
    for name in xlsx_names:
        wb = openpyxl.load_workbook(os.path.join(SRC, name), data_only=True)
        sheets.extend(wb.worksheets)

    analyses = [analyze_sheet(ws) for ws in sheets]
    num_to_index = {}
    for i, ws in enumerate(sheets):
        m = re.match(r"^(\d+)\.", ws.title)
        if m:
            num_to_index[int(m.group(1))] = i
    merged_indices = {i for i, a in enumerate(analyses) if tuple(a["header"][:2]) == ("分支", "序列／前提")}

    tabs = []
    panels = []
    for i, ws in enumerate(sheets):
        sid = "sheet-%d" % (i + 1)
        active = " is-active" if i == 0 else ""
        tabs.append('<button class="sheet-tab%s" data-target="%s">%s</button>' %
                     (active, sid, html.escape(ws.title)))
        content = render_sheet_panel(analyses[i], i, num_to_index, merged_indices)
        panels.append('<div class="sheet-panel%s" id="%s">%s</div>' % (active, sid, content))

    body = page_head(
        eyebrow, title,
        ["%d 個工作表" % len(sheets), "來源：%s" % "、".join(xlsx_names)],
        lede,
        crumb_html='<div class="crumb"><a href="../index.html">首頁</a> ／ 叫序資料 ／ %s</div>' % title,
    )
    body += '<section>'
    body += '<div class="sheet-tabs" data-sheet-tabs="#%s-panels">%s</div>' % (title, "".join(tabs))
    body += '<div id="%s-panels">%s</div>' % (title, "".join(panels))
    body += '</section>'
    body += '<footer class="site-footer">資料整理自 %s；點數為由 rp 換算的大牌點（HCP），原文 rp 保留於括號。</footer>' % html.escape(SOURCE_BOOK)

    html_out = render_page(
        title="%s ｜ 叫序資料 ｜ %s" % (title, SITE_NAME),
        description="%s 的完整叫序資料表" % title,
        root_prefix="../",
        current_href="叫序資料/%s.html" % title,
        body_html=body,
    )
    write_file(os.path.join(OUT, "叫序資料", "%s.html" % title), html_out)




# --------------------------------------------------------------------------
# 首頁
# --------------------------------------------------------------------------

def count_sheets(xlsx_name):
    wb = openpyxl.load_workbook(os.path.join(SRC, xlsx_name), read_only=True)
    return len(wb.sheetnames)


def build_index():
    body = page_head(
        "Terrorist's Moscito",
        SITE_NAME,
        ["%d 份叫序資料表" % len(PAGES), "點數：HCP（附原文 rp）"],
        "整理自 Bo-Yin Yang《Terrorist's Moscito》（Major-Oriented Strong Club, with Intrepid Two Openers）。"
        "收錄 1♣、1♦、高花、1NT、阻擊五種開叫的叫序表，以及接力方案與競叫兩份通用模塊。"
        "原文以 rp（A=3、K=2、Q=1）計算強度，本站換算成大牌點，原文 rp 保留在括號內。",
    )

    cards = []
    for xlsx_name, title, kicker, lede in PAGES:
        cards.append(
            '<a class="doc-card" href="叫序資料/%s.html"><div class="k">%s</div><h3>%s</h3>'
            '<p>%s共 %d 個工作表。</p></a>' % (
                title, html.escape(kicker), html.escape(title),
                suit_colorize(lede), count_sheets(xlsx_name)))

    body += """
<section>
  <div class="sec-head"><h2>開叫叫序</h2><span class="tag">依開叫整理的叫序表</span></div>
  <div class="card-grid">%s</div>
</section>

<section>
  <div class="sec-head"><h2>通用模塊</h2><span class="tag">各開叫共用的接力與競叫結構</span></div>
  <div class="card-grid">%s</div>
</section>

<div class="external-card">
  <p><strong>rp → 大牌點換算</strong> — 以 40 萬副隨機牌模擬：大牌點 ≈ 1.4 × rp + 1.5；每份叫序表的「說明」分頁最後都附完整換算表。</p>
  <a href="叫序資料/接力方案.html">看接力方案</a>
</div>
""" % ("".join(cards[:5]), "".join(cards[5:]))
    body += ('<footer class="site-footer">整理自 %s。以上資料僅供個人叫牌系統學習與對照使用。</footer>'
             % html.escape(SOURCE_BOOK))

    html_out = render_page(
        title=SITE_NAME,
        description="整理自 Terrorist's Moscito 的叫序資料庫（HCP 版）",
        root_prefix="",
        current_href=None,
        body_html=body,
    )
    write_file(os.path.join(OUT, "index.html"), html_out)


# --------------------------------------------------------------------------

def main():
    build_index()
    for xlsx_name, title, _kicker, lede in PAGES:
        build_xlsx_page(xlsx_name, title, "叫序資料 · %s" % title, suit_colorize(lede))


if __name__ == "__main__":
    main()
